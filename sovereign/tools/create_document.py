"""
create_document() agent tool.

Milestone 4 (feat/docx-generation): a structured Approval Note template
— title, generated-on line, optional source-documents line, then
ordered sections with heading/body and an optional evidence sub-list
(source, location, quoted text) rendered distinctly from the narrative
body, satisfying the brief's "preserve source/evidence references"
requirement. Backward compatible with Milestone 3's plain
{"heading","body"} sections — "evidence" and source_documents are both
optional, so existing callers are unaffected.
"""

from datetime import datetime
from pathlib import Path

import docx


class DocumentCreationError(Exception):
    pass


def _add_evidence(doc, evidence):
    doc.add_heading("Evidence", level=2)
    for item in evidence:
        text = (item.get("text") or "").strip()
        if not text:
            continue
        source = (item.get("source") or "unknown source").strip()
        location = (item.get("location") or "").strip()
        label = f"{source} ({location})" if location else source

        p = doc.add_paragraph(style="List Bullet")
        bold_run = p.add_run(f"{label}: ")
        bold_run.bold = True
        quote_run = p.add_run(f'"{text}"')
        quote_run.italic = True


def create_document(title: str, sections, output_dir: str, source_documents=None) -> str:
    """
    sections: list of dicts, each —
        heading: str
        body: str
        evidence: optional list of {"source": str, "location": str, "text": str}
    source_documents: optional list of filenames, rendered as a short
        traceability line under the title.
    """
    title = (title or "").strip()
    if not title:
        raise DocumentCreationError("A title is required")
    if not sections:
        raise DocumentCreationError("At least one section is required")

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_name = "".join(c if c.isalnum() or c in " -_" else "_" for c in title).strip()
    safe_name = safe_name or "document"
    out_path = out_dir / f"{safe_name}.docx"

    d = docx.Document()
    d.add_heading(title, level=0)

    meta = d.add_paragraph()
    meta.add_run(f"Generated {datetime.now():%Y-%m-%d %H:%M}").italic = True

    if source_documents:
        src = d.add_paragraph()
        src.add_run("Source documents: " + ", ".join(source_documents)).italic = True

    written_sections = 0
    for section in sections:
        heading = (section.get("heading") or "").strip()
        body = (section.get("body") or "").strip()
        evidence = section.get("evidence") or []
        if not heading and not body and not evidence:
            continue
        if heading:
            d.add_heading(heading, level=1)
        if body:
            d.add_paragraph(body)
        if evidence:
            _add_evidence(d, evidence)
        written_sections += 1

    if written_sections == 0:
        raise DocumentCreationError("Every section was empty — nothing to write")

    d.save(str(out_path))
    return f"Created {out_path.name} with {written_sections} section(s) at {out_path}"
