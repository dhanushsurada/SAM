"""
Offline smoke test for Sovereign Workbench document ingestion (SIH26117,
Milestone 1) — sovereign/ingestion/{extract,chunk,pipeline}.py.

PDF and DOCX fixtures are generated on the fly with reportlab/python-docx
so no binary fixtures are committed to the repo. Only the image-extraction
path touches the network (Ollama's vision endpoint), and that's mocked
here exactly like the existing tests mock requests.post — nothing in
this file needs a live Ollama or Mac.

Usage:
    HOME=/tmp/sam_smoke_sovereign_ingestion python3 tests/test_sovereign_ingestion_offline.py
"""

import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).parent.parent))

from sovereign.ingestion import (  # noqa: E402
    Chunk,
    ExtractionError,
    chunk_document,
    chunk_text,
    detect_file_type,
    extract,
    ingest_file,
    normalize,
)

results = []


def check(label, condition):
    status = "PASS" if condition else "FAIL"
    results.append(condition)
    print(f"[{status}] {label}")


TMP = Path(tempfile.mkdtemp(prefix="sam_sovereign_test_"))


def _make_pdf(path: Path, pages_text):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    for i, text in enumerate(pages_text):
        c.drawString(72, 720, text)
        if i < len(pages_text) - 1:
            c.showPage()  # only between pages — save() finalizes the last one
    c.save()


def _make_docx(path: Path):
    import docx
    d = docx.Document()
    d.add_heading("Inspection Checklist", level=1)
    d.add_paragraph("All valves must be inspected annually.")
    d.add_paragraph("Pressure readings must not exceed 150 PSI.")
    d.add_heading("Approval Requirements", level=1)
    d.add_paragraph("Two signatures are required for approval.")
    d.save(str(path))


def test_detect_file_type():
    check("Detects .pdf", detect_file_type(Path("x.pdf")) == "pdf")
    check("Detects .docx", detect_file_type(Path("x.docx")) == "docx")
    check("Detects .txt", detect_file_type(Path("x.txt")) == "txt")
    check("Detects .png as image", detect_file_type(Path("x.png")) == "image")
    try:
        detect_file_type(Path("x.exe"))
        check("Unsupported extension raises ExtractionError", False)
    except ExtractionError:
        check("Unsupported extension raises ExtractionError", True)


def test_extract_pdf():
    pdf_path = TMP / "inspection_report.pdf"
    _make_pdf(pdf_path, ["Page one: boiler pressure log.", "Page two: signed off by inspector."])
    doc = extract(str(pdf_path))
    check("PDF extracted with 2 pages", len(doc.pages) == 2)
    check("PDF page 1 text captured", "boiler pressure" in doc.pages[0].text.lower())
    check("PDF page 2 text captured", "signed off" in doc.pages[1].text.lower())
    check("PDF page labels are 'page N'", doc.pages[0].label == "page 1" and doc.pages[1].label == "page 2")
    check("PDF file_type is 'pdf'", doc.file_type == "pdf")
    check("PDF content_hash is a 64-char sha256 hex digest", len(doc.content_hash) == 64)


def test_extract_docx_sections_from_headings():
    docx_path = TMP / "sop.docx"
    _make_docx(docx_path)
    doc = extract(str(docx_path))
    labels = [p.label for p in doc.pages]
    check("DOCX sections derived from headings",
          labels == ["Inspection Checklist", "Approval Requirements"])
    check("DOCX first section body captured", "150 PSI" in doc.pages[0].text)
    check("DOCX second section body captured", "Two signatures" in doc.pages[1].text)


def test_extract_txt():
    txt_path = TMP / "notes.txt"
    txt_path.write_text("Deviation found in section 4.2.\nRequires recalculation.")
    doc = extract(str(txt_path))
    check("TXT extracted as a single section", len(doc.pages) == 1)
    check("TXT text captured", "Deviation found" in doc.pages[0].text)


def test_extract_image_mocked():
    img_path = TMP / "scan.png"
    from PIL import Image
    Image.new("RGB", (10, 10), color="white").save(img_path)

    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.raise_for_status = lambda: None
    mock_resp.json.return_value = {"response": "Reading: 148 PSI"}
    with patch("sovereign.ingestion.extract.requests.post", return_value=mock_resp) as mock_post:
        doc = extract(str(img_path), ollama_host="http://localhost:11434", vision_model="moondream")
        check("Image extraction calls Ollama's vision endpoint", mock_post.called)
        check("Image extraction returns the mocked transcription", doc.pages[0].text == "Reading: 148 PSI")
        check("Image section labeled 'image'", doc.pages[0].label == "image")


def test_extract_image_model_failure_raises():
    img_path = TMP / "scan2.png"
    from PIL import Image
    Image.new("RGB", (10, 10), color="white").save(img_path)
    with patch("sovereign.ingestion.extract.requests.post", side_effect=ConnectionError("no route")):
        try:
            extract(str(img_path))
            check("Vision model failure raises ExtractionError (not silent)", False)
        except ExtractionError:
            check("Vision model failure raises ExtractionError (not silent)", True)


def test_missing_file_raises():
    try:
        extract(str(TMP / "does_not_exist.pdf"))
        check("Missing file raises FileNotFoundError", False)
    except FileNotFoundError:
        check("Missing file raises FileNotFoundError", True)


def test_directory_raises():
    try:
        extract(str(TMP))
        check("Passing a directory raises IsADirectoryError", False)
    except IsADirectoryError:
        check("Passing a directory raises IsADirectoryError", True)


def test_empty_file_raises():
    empty_path = TMP / "empty.txt"
    empty_path.write_bytes(b"")
    try:
        extract(str(empty_path))
        check("Empty file raises ExtractionError", False)
    except ExtractionError:
        check("Empty file raises ExtractionError", True)


def test_corrupted_pdf_raises_cleanly():
    fake_pdf = TMP / "corrupted.pdf"
    fake_pdf.write_bytes(b"this is not a real pdf file, just garbage bytes")
    try:
        extract(str(fake_pdf))
        check("Corrupted PDF raises ExtractionError (not a raw traceback)", False)
    except ExtractionError:
        check("Corrupted PDF raises ExtractionError (not a raw traceback)", True)


def test_normalize_collapses_whitespace():
    messy = "Too    many   spaces\r\n\r\n\r\n\r\nand blank lines"
    clean = normalize(messy)
    check("normalize() collapses runs of spaces", "Too many spaces" in clean)
    check("normalize() caps blank-line runs", "\n\n\n" not in clean)


def test_chunk_text_basic():
    text = " ".join(f"word{i}" for i in range(1000))
    chunks = chunk_text(text, chunk_size=300, chunk_overlap=50)
    check("chunk_text produces multiple chunks for long text", len(chunks) > 1)
    check("Every chunk is non-empty", all(c.strip() for c in chunks))
    check("Consecutive chunks overlap correctly", chunks[0].split()[-1] == chunks[1].split()[49])


def test_chunk_text_edge_cases():
    check("Empty text yields no chunks", chunk_text("") == [])
    check("Whitespace-only text yields no chunks", chunk_text("   \n  ") == [])
    short = "just a few words here"
    check("Short text yields exactly one chunk",
          chunk_text(short, chunk_size=300, chunk_overlap=50) == [short])
    try:
        chunk_text("a b c", chunk_size=10, chunk_overlap=10)
        check("overlap >= chunk_size raises ValueError", False)
    except ValueError:
        check("overlap >= chunk_size raises ValueError", True)


def test_chunk_document_metadata_and_provenance():
    pdf_path = TMP / "metadata_check.pdf"
    _make_pdf(pdf_path, ["Short first page.", "Short second page."])
    doc = extract(str(pdf_path))
    chunks = chunk_document(doc, chunk_size=300, chunk_overlap=50)
    check("chunk_document produces one chunk per short page", len(chunks) == 2)
    check("Chunk filename preserved", chunks[0].filename == "metadata_check.pdf")
    check("Chunk section label preserved from page", chunks[0].section == "page 1")
    check("Chunk section_index preserved", chunks[1].section_index == 2)
    check("Chunk content_hash matches document", chunks[0].content_hash == doc.content_hash)
    check("chunk_id format is doc_id:section_index:chunk_index",
          chunks[0].chunk_id == f"{doc.content_hash[:16]}:1:0")


def test_doc_id_stable_across_reingestion():
    pdf_path = TMP / "stability_check.pdf"
    _make_pdf(pdf_path, ["Same content every time."])
    doc1 = extract(str(pdf_path))
    doc2 = extract(str(pdf_path))
    check("Re-extracting identical bytes yields the same content_hash",
          doc1.content_hash == doc2.content_hash)
    chunks1, chunks2 = chunk_document(doc1), chunk_document(doc2)
    check("Re-ingesting identical bytes yields the same doc_id",
          chunks1[0].doc_id == chunks2[0].doc_id)
    check("Re-ingesting identical bytes yields the same chunk_id",
          chunks1[0].chunk_id == chunks2[0].chunk_id)


def test_ingest_file_end_to_end():
    pdf_path = TMP / "e2e.pdf"
    _make_pdf(pdf_path, ["End to end ingestion pipeline test content."])
    chunks = ingest_file(str(pdf_path))
    check("ingest_file returns at least one chunk", len(chunks) >= 1)
    check("ingest_file chunk is a Chunk instance", isinstance(chunks[0], Chunk))
    check("ingest_file preserves filename", chunks[0].filename == "e2e.pdf")


def main():
    test_detect_file_type()
    test_extract_pdf()
    test_extract_docx_sections_from_headings()
    test_extract_txt()
    test_extract_image_mocked()
    test_extract_image_model_failure_raises()
    test_missing_file_raises()
    test_directory_raises()
    test_empty_file_raises()
    test_corrupted_pdf_raises_cleanly()
    test_normalize_collapses_whitespace()
    test_chunk_text_basic()
    test_chunk_text_edge_cases()
    test_chunk_document_metadata_and_provenance()
    test_doc_id_stable_across_reingestion()
    test_ingest_file_end_to_end()

    print(f"\n{sum(results)}/{len(results)} checks passed.")
    if not all(results):
        sys.exit(1)
    print("Sovereign Workbench document ingestion (Milestone 1) verified.")


if __name__ == "__main__":
    main()
